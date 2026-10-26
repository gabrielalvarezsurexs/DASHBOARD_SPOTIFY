"""Informe ejecutivo en Word (out/informe.docx), con Montserrat embebida.

Uso: uv run python -m pipeline.informe_docx   (requiere informe/fig, ver pipeline.eda_figures)

Reglas de formato que pidió el equipo: todo en negro, Montserrat, títulos discretos, sin rayas (em
dash). Las cifras son de la ventana común y coinciden con el dashboard en «Ventana común».
La fuente se embebe para que el documento se vea igual en una computadora sin Montserrat.
"""

import re
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pipeline.config import ROOT

FIG = ROOT / "informe" / "fig"
FONTS = ROOT / "informe" / "fonts"
OUT = ROOT / "out" / "informe.docx"
FONT = "Montserrat"
BLACK = RGBColor(0, 0, 0)
BODY_PT = 9.5


# ---------- estilos ----------
def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    for name, size, bold, before, after in (
        ("Normal", BODY_PT, False, 0, 4),
        ("Title", 15, True, 0, 2),
        ("Heading 1", 11.5, True, 10, 3),
        ("Heading 2", BODY_PT, True, 6, 1),
        ("Caption", 8, False, 1, 6),
    ):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = False
        st.font.color.rgb = BLACK
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(attr), None)
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.15
        if name.startswith("Heading"):
            pf.keep_with_next = True
        if name == "Title":
            # el estilo Title de la plantilla trae un borde inferior azul
            ppr = st.element.get_or_add_pPr()
            for b in ppr.findall(qn("w:pBdr")):
                ppr.remove(b)
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.space_after = Pt(2)
        doc.styles[name].paragraph_format.line_spacing = 1.15


INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def runs(p, text: str) -> None:
    """**negrita** y *cursiva* dentro del texto."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def para(doc, text: str, style: str | None = None):
    p = doc.add_paragraph(style=style)
    runs(p, text)
    return p


def bullets(doc, items: list[str], numbered: bool = False) -> None:
    for it in items:
        para(doc, it, "List Number" if numbered else "List Bullet")


def borders(table, inner: bool = True) -> None:
    """Solo líneas horizontales finas: encabezado, filas y cierre."""
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        on = edge in ("top", "bottom") or (edge == "insideH" and inner)
        e.set(qn("w:val"), "single" if on else "nil")
        if on:
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
        b.append(e)
    tbl_pr.append(b)


def table(doc, header: list[str], rows: list[list[str]], widths: list[float], right=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    borders(t)
    for r, values in enumerate([header, *rows]):
        for c, v in enumerate(values):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if c in right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(v)
            run.font.size = Pt(8.5)
            run.bold = r == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def figure(doc, png: str, width_cm: float, caption: str) -> None:
    doc.add_picture(str(FIG / png), width=Cm(width_cm))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(doc, caption, "Caption")


def fig_grid(doc, items: list[tuple[str, str]]) -> None:
    """Figuras del EDA en una cuadrícula 2×2 sin bordes."""
    t = doc.add_table(rows=2, cols=2)
    borders(t, inner=False)
    tbl_pr = t._tbl.tblPr
    for b in tbl_pr.findall(qn("w:tblBorders")):
        for e in b:
            e.set(qn("w:val"), "nil")
    for i, (png, cap) in enumerate(items):
        cell = t.cell(i // 2, i % 2)
        cell.width = Cm(8.75)
        p = cell.paragraphs[0]
        p.add_run().add_picture(str(FIG / png), width=Cm(8.3))
        c = cell.add_paragraph(style="Caption")
        runs(c, cap)


# ---------- contenido ----------
def content(doc: Document) -> None:
    doc.add_paragraph("Cómo escuchamos: tres años de Spotify de Gabriel y Javier", "Title")
    para(
        doc,
        "Informe ejecutivo · Dashboard e insights de la red social · Big Data · "
        "Gabriel y Javier · 25 de octubre de 2026",
    ).runs[0].font.size = Pt(8.5)

    doc.add_heading("1. Resumen ejecutivo", 1)
    para(
        doc,
        "Pedimos a Spotify nuestro historial extendido (*Extended Streaming History*), lo "
        "limpiamos con un pipeline reproducible y construimos un solo dashboard interactivo que "
        "compara a los dos integrantes. En total son 81,418 reproducciones. Para que la "
        "comparación sea justa usamos la ventana en la que ambos tenemos datos continuos: de julio "
        "de 2023 a agosto de 2026. Respondimos las preguntas del proyecto: cuándo somos más "
        "activos, cuáles son nuestras interacciones más densas y cuáles las menos.",
    )
    table(
        doc,
        ["Ventana común (jul 2023 a ago 2026)", "Gabriel", "Javier"],
        [
            ["Horas escuchadas", "2,841", "713"],
            ["Días con escucha", "95% (1,087 días)", "49% (546 días)"],
            ["Mediana por día con escucha", "139 min", "63 min"],
            ["Hora pico", "14 h (y otro a las 6 h)", "18 h"],
            ["Artistas distintos", "1,662", "1,058"],
            ["Peso de sus 10 artistas principales", "21% del tiempo", "28% del tiempo"],
            ["Canciones saltadas (skips)", "21%", "15%"],
            ["Dispositivo principal", "Android (92%)", "Windows (70%)"],
            ["Tiempo en podcasts", "0.1%", "5.2%"],
        ],
        [7.6, 4.8, 4.8],
        right=(1, 2),
    )
    bullets(
        doc,
        [
            "**Intensidad.** Gabriel escucha casi todos los días y unas 4 veces más horas que "
            "Javier. Javier escucha menos de la mitad de los días, pero cuando escucha lo hace "
            "alrededor de una hora.",
            "**Horario.** Para Gabriel la música acompaña todo el día hábil desde las 6 de la "
            "mañana. Javier concentra el 40% de su tiempo entre las 17 y las 21 h y casi nunca "
            "escucha de madrugada (0.4%).",
            "**Tendencia.** Gabriel pasó de unas 42 horas al mes en 2023 a 104 en 2025. Javier "
            "bajó a 10 horas al mes en 2025 y repuntó en 2026 durante una estancia en Estados "
            "Unidos.",
            "**Lealtad.** Bad Bunny fue el artista principal de Gabriel seis trimestres seguidos. "
            "Alfa, un cantante italiano, lo fue de Javier en 8 de los últimos 9 trimestres.",
            "**En común.** Compartimos 205 artistas. De los que ambos descubrimos dentro de la "
            "ventana, Javier llegó primero a 55 y Gabriel a 27.",
        ],
    )

    doc.add_heading("2. Metodología", 1)
    doc.add_heading("Datos", 2)
    para(
        doc,
        "Cada uno solicitó el historial extendido desde la configuración de privacidad de Spotify. "
        "El export trae varios JSON con una fila por reproducción y 23 campos: fecha y hora de fin "
        "(ts, en UTC), milisegundos reproducidos, canción, artista, álbum, podcast, audiolibro, "
        "plataforma, país de conexión, IP, shuffle, skip, motivo de inicio y de fin, offline e "
        "incógnito. "
        "Gabriel aporta 61,422 reproducciones y Javier 19,996.",
    )
    doc.add_heading("Pipeline", 2)
    bullets(
        doc,
        [
            "**Raw.** Los zips se descomprimen tal cual en data/raw/<persona>/. La persona sale "
            "del nombre de la carpeta, así que sumar a alguien más es soltar su zip.",
            "**Silver (DuckDB a Parquet).** Tipado, hora de inicio (ts menos ms_played), hora "
            "local con la zona horaria del país de conexión, bandera de stream, familia de "
            "plataforma. Se quitan 11 duplicados exactos y la columna de IP.",
            "**Dashboard.** Un script exporta la tabla limpia a un dashboard estático (HTML con "
            "Observable Plot) que filtra en el navegador por persona, periodo y tipo de "
            "reproducción. Se abre con doble clic, sin servidor.",
        ],
        numbered=True,
    )
    para(
        doc,
        "La infraestructura del curso (MinIO, Spark, ClickHouse y Kafka en Docker) quedó montada "
        "y probada, pero con 81 mil filas (unos 60 MB) DuckDB procesa todo en menos de 2 segundos "
        "en una laptop. Elegimos la herramienta por el volumen real; Spark y ClickHouse quedan "
        "como ruta de escalado si se suman más personas. Todo se regenera con un solo comando "
        "(generar.cmd).",
    )
    doc.add_heading("Reglas de limpieza y cómo traducimos las preguntas", 2)
    table(
        doc,
        ["Pregunta del proyecto", "Cómo la medimos en Spotify"],
        [
            ["¿Cuándo soy más activo?", "Minutos por hora, día de la semana y mes, en hora local"],
            ["Interacciones más densas", "Artistas y canciones por tiempo; días récord; eras"],
            ["Interacciones menos densas", "Skips: qué se abandona, a qué hora y en qué modo"],
            ["¿Qué tanto nos parecemos?", "Artistas en común y quién llegó primero"],
        ],
        [6.0, 11.2],
    )
    bullets(
        doc,
        [
            "**Stream:** reproducción de 30 s o más. Las más cortas solo cuentan para medir skips.",
            "**Hora local:** por país de conexión (MX a America/Mexico_City), nunca con un "
            "desfase fijo.",
            "**Skip:** el campo skipped de Spotify, tal cual.",
            "**Comparaciones:** en porcentaje del tiempo propio y dentro de la ventana común.",
        ],
    )

    doc.add_heading("3. Principales hallazgos", 1)
    doc.add_heading("Para Gabriel la música es rutina diaria; para Javier, un rato por la tarde", 2)
    figure(
        doc,
        "hallazgo_heatmap.png",
        13.0,
        "Figura 1. Porcentaje del tiempo de cada persona por hora y día de la semana. Cada mapa "
        "usa su propia escala para que el volumen de Gabriel no tape el patrón de Javier.",
    )
    para(
        doc,
        "Gabriel tiene un pico a las 6 h, que coincide con el trayecto de la mañana, y escucha de "
        "forma pareja hasta la noche; el fin de semana pesa menos (20% del tiempo contra el 29% "
        "esperado). Javier escucha casi solo de 17 a 21 h y su pico más fuerte es el domingo a "
        "las 18 h. El dispositivo lo confirma: Gabriel escucha en el teléfono (92% Android), en "
        "movimiento; Javier en la computadora (70% Windows), frente al escritorio.",
    )
    doc.add_heading("Uno aceleró y el otro tuvo un año en pausa", 2)
    figure(
        doc,
        "hallazgo_mensual.png",
        13.0,
        "Figura 2. Horas por mes de cada persona en la ventana común. Agosto de 2026 no se "
        "grafica porque está incompleto: los datos terminan el 3 (Javier) y el 25 (Gabriel).",
    )
    para(
        doc,
        "Gabriel más que duplicó su escucha entre 2023 y 2025, con un día récord de 8.8 horas el "
        "4 de diciembre de 2025. Javier bajó a 10 horas al mes en 2025 y repuntó en mayo de 2026: "
        "su día récord (6 horas) fue el 2 de junio de 2026, al final de su estancia en Estados "
        "Unidos.",
    )
    doc.add_heading("Interacciones más densas: lealtad contra variedad", 2)
    para(
        doc,
        "La canción más repetida de Gabriel es «BAILE INoLVIDABLE» de Bad Bunny, con 412 streams, "
        "y Bad Bunny ocupa el 6.2% de su tiempo; antes de 2025 su artista principal cambiaba cada "
        "trimestre. Javier repite a Alfa («CI SARÒ», 74 streams), que ocupa el 5.3% de su "
        "tiempo, y a Irama («Tu No», 70). Además dedica el 5% de su escucha a podcasts, sobre "
        "todo uno en francés (23 horas). Javier conoce menos artistas, pero vuelve más a ellos: "
        "sus 10 principales suman el 28% de su tiempo contra el 21% de Gabriel.",
    )
    doc.add_heading("Interacciones menos densas: lo que saltamos", 2)
    para(
        doc,
        "Gabriel salta el 27% de las canciones de noche y el 16% en la mañana. Contra lo "
        "esperado, salta más cuando no usa shuffle (31% contra 19%): cuando elige qué escuchar, "
        "busca activamente hasta dar con la canción. Javier salta alrededor del 15% sin importar "
        "la hora ni el modo.",
    )
    doc.add_heading("Compartimos 205 artistas, pero casi nunca al mismo tiempo", 2)
    para(
        doc,
        "Sebastian Yatra, Reik, Alex Ponce y Coldplay son los que más pesan para los dos. Los "
        "artistas compartidos ocupan el 41% del tiempo de Gabriel y el 42% del de Javier. En los "
        "82 artistas que los dos empezamos a escuchar dentro de la ventana, con más de 30 días "
        "de diferencia, Javier llegó primero a 55 y Gabriel a 27. Tiene un sesgo que reportamos: "
        "el historial continuo de Javier empieza en noviembre de 2022, ocho meses antes que el "
        "de Gabriel, así que algunos artistas que parecen nuevos para él quizá ya los conocía.",
    )

    doc.add_heading("4. Visualizaciones clave del EDA", 1)
    para(
        doc,
        "Estas gráficas no están en el dashboard: son las que nos hicieron cambiar la limpieza.",
    )
    fig_grid(
        doc,
        [
            (
                "eda_cobertura.png",
                "**Cobertura.** Al rellenar los meses vacíos con cero apareció el hueco: los datos "
                "de Gabriel empiezan el 14 de julio de 2023. Con una ventana desde enero, en 2023 "
                "Javier parecía escuchar más que Gabriel (0.9 a 1); en los meses que ambos tienen "
                "datos, Gabriel escuchó el doble. Movimos la ventana a julio.",
            ),
            (
                "eda_zona_horaria.png",
                "**Zona horaria.** ts está en UTC y marca el fin de la reproducción. Sin "
                "convertir, el pico de Javier cae a medianoche. DuckDB además leía las fechas "
                "con la zona de la computadora, así que fijamos la sesión en UTC.",
            ),
            (
                "eda_duracion.png",
                "**Duración.** Hay dos grupos: saltadas en 0 a 5 s (10% de las reproducciones de "
                "Gabriel) y canciones completas de unos 3 minutos. La regla de 30 s los separa.",
            ),
            (
                "eda_skip.png",
                "**Skip.** Dudábamos entre skipped y reason_end = fwdbtn. skipped vale 97 a 98% "
                "en los botones de saltar y 0% en el resto, así que usamos el campo directo.",
            ),
        ],
    )

    doc.add_heading("5. Reflexiones sobre privacidad digital", 1)
    bullets(
        doc,
        [
            "**Cada reproducción trae la IP.** Son 81 mil puntos de ubicación aproximada. La "
            "eliminamos en el primer paso del pipeline y no llega a ninguna gráfica.",
            "**El país de conexión reconstruye viajes con fecha exacta.** Aun sin la IP se ve que "
            "Gabriel estuvo en España, Reino Unido, Francia, Chequia, Alemania e Italia entre el "
            "18 de junio y el 9 de julio de 2025, y que Javier se conectó desde Estados Unidos "
            "entre el 30 de marzo y el 2 de junio de 2026.",
            "**El horario revela la rutina.** A qué hora sale alguien de casa, qué días no "
            "trabaja y cuándo duerme se lee en un heatmap. Junto con los viajes, es información "
            "útil para cualquiera con malas intenciones.",
            "**Privado no significa no registrado.** Las 2,073 reproducciones de Gabriel en modo "
            "incógnito desde noviembre de 2024 aparecen completas en el export. El modo privado "
            "solo las oculta de los amigos y de las recomendaciones, no de Spotify.",
            "**Qué decidimos publicar.** Los zips originales no se suben al repositorio. El "
            "dashboard usa la tabla limpia, sin IP, y ambos revisamos nuestros datos antes de "
            "aceptar mostrarlos.",
        ],
    )

    doc.add_heading("6. Conclusiones y aprendizajes", 1)
    bullets(
        doc,
        [
            "**Lo que dicen los datos.** Escuchamos música de forma casi opuesta. Para Gabriel es "
            "el fondo de todo el día, con mucha variedad y muchos skips cuando elige. Para Javier "
            "es un momento de la tarde, con pocos artistas a los que vuelve y podcasts en otros "
            "idiomas. Cada afirmación se puede comprobar en el dashboard con los filtros.",
            "**Lo técnico.** La mayor parte del trabajo fue entender el export, no graficar: ts "
            "marca el fin y no el inicio, está en UTC, hay meses vacíos que una línea esconde y "
            "campos que parecen redundantes. Cada regla de limpieza nació de una gráfica "
            "exploratoria.",
            "**Lo metodológico.** Comparar a dos personas exige una ventana común y medidas "
            "relativas al tiempo de cada uno. Si no, siempre gana quien tiene más datos.",
            "**Lo personal.** Ver la propia rutina dibujada en un heatmap es más revelador, e "
            "incómodo, que cualquier resumen anual de Spotify.",
        ],
    )
    p = para(
        doc,
        "El proceso completo (obstáculos, dudas, decisiones y aprendizajes) está en la bitácora. "
        "El dashboard se abre con dashboard/web/index.html y todo se regenera con generar.cmd.",
    )
    p.runs[0].font.size = Pt(8)


# ---------- fuente embebida ----------
def _obfuscate(ttf: bytes, key: str) -> bytes:
    """ECMA-376 §17.8.1: XOR de los primeros 32 bytes con la GUID en orden inverso."""
    k = bytes.fromhex(key.replace("-", ""))[::-1]
    head = bytes(b ^ k[i % 16] for i, b in enumerate(ttf[:32]))
    return head + ttf[32:]


def embed_fonts(docx: Path) -> None:
    faces = {"Regular": "embedRegular", "Bold": "embedBold", "Italic": "embedItalic"}
    with zipfile.ZipFile(docx) as z:
        files = {n: z.read(n) for n in z.namelist()}

    rels, embeds = [], []
    for i, (face, tag) in enumerate(faces.items(), start=1):
        key = str(uuid.uuid4()).upper()
        files[f"word/fonts/font{i}.odttf"] = _obfuscate(
            (FONTS / f"Montserrat-{face}.ttf").read_bytes(), key
        )
        rels.append(
            f'<Relationship Id="rIdF{i}" Target="fonts/font{i}.odttf" Type="http://schemas.'
            'openxmlformats.org/officeDocument/2006/relationships/font"/>'
        )
        embeds.append(f'<w:{tag} r:id="rIdF{i}" w:fontKey="{{{key}}}"/>')
    files["word/_rels/fontTable.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rels)
        + "</Relationships>"
    ).encode()

    ft = files["word/fontTable.xml"].decode()
    if 'xmlns:r="' not in ft.split(">", 2)[1]:
        ft = ft.replace(
            "<w:fonts ",
            '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships" ',
            1,
        )
    ft = re.sub(r'<w:font w:name="Montserrat">.*?</w:font>', "", ft, flags=re.S)
    ft = ft.replace("</w:fonts>", f'<w:font w:name="{FONT}">{"".join(embeds)}</w:font></w:fonts>')
    files["word/fontTable.xml"] = ft.encode()

    ct = files["[Content_Types].xml"].decode()
    if 'Extension="odttf"' not in ct:
        ct = ct.replace(
            "<Default ",
            '<Default Extension="odttf" '
            'ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ',
            1,
        )
    files["[Content_Types].xml"] = ct.encode()

    st = files["word/settings.xml"].decode()
    if "embedTrueTypeFonts" not in st:
        st = re.sub(r"(<w:zoom[^>]*/>)", r"\1<w:embedTrueTypeFonts/>", st, count=1)
    files["word/settings.xml"] = st.encode()

    with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)


def build(out: Path = OUT) -> Path:
    doc = Document()
    setup(doc)
    content(doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    embed_fonts(out)
    return out


if __name__ == "__main__":
    print(build())
