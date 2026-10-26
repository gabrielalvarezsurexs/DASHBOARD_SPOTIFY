"""Auditoría de cifras: recalcula desde los datos cada número que aparece en el informe.

Uso: uv run python -m pipeline.auditoria
Salida: docs/auditoria-cifras.md (ID, afirmación, cita, valor esperado, valor calculado, estado)

Cada fila es una afirmación del informe (out/informe.docx, texto en pipeline/informe_docx.py) con
el valor tal como está escrito. El valor calculado se formatea con el mismo redondeo que usa el
informe, así que «OK» significa que el texto publicado coincide con los datos. Las reglas de cada
cálculo replican las del dashboard (dashboard/web/app.js); ver docs/tecnico.md §6.
"""

import json
import re
from collections.abc import Callable
from dataclasses import dataclass

import duckdb

from pipeline.config import COMMON_WINDOW_START as W
from pipeline.config import RAW_DIR, ROOT, SILVER_PATH

OUT = ROOT / "docs" / "auditoria-cifras.md"
DATA_JS = ROOT / "dashboard" / "web" / "data.js"
INFORME = ROOT / "out" / "informe.docx"
RAW_GLOB = (RAW_DIR / "*" / "*" / "Streaming_History_*.json").as_posix()

# Texto literal del informe que respalda cada afirmación: se busca en out/informe.docx
CITAS = {
    "A01": "81,418 reproducciones",
    "A02": "Gabriel aporta 61,422 reproducciones",
    "A03": "Javier 19,996",
    "A04": "Se quitan 11 duplicados",
    "A05": "23 campos",
    "A06": "Son 81 mil puntos",
    "A07": "La eliminamos en el primer paso",
    "B01": "2,841",
    "B02": "713",
    "B03": "unas 4 veces más horas",
    "B04": "95% (1,087 días)",
    "B05": "49% (546 días)",
    "B06": "139 min",
    "B07": "63 min",
    "B08": "14 h (y otro a las 6 h)",
    "B09": "18 h",
    "B10": "y otro a las 6 h",
    "B11": "1,662",
    "B12": "1,058",
    "B13": "21% del tiempo",
    "B14": "28% del tiempo",
    "B15": "21%",
    "B16": "15%",
    "B17": "Android (92%)",
    "B18": "Windows (70%)",
    "B19": "0.1%",
    "B20": "5.2%",
    "C01": "concentra el 40% de su tiempo entre las 17 y las 21 h",
    "C02": "(0.4%)",
    "C03": "unas 42 horas al mes en 2023",
    "C04": "a 104 en 2025",
    "C05": "bajó a 10 horas al mes en 2025",
    "C06": "seis trimestres seguidos",
    "C07": "8 de los últimos 9 trimestres",
    "C08": "Javier llegó primero a 55 y Gabriel a 27",
    "D01": "(20% del tiempo contra el 29% esperado)",
    "D02": "el domingo a las 18 h",
    "D03": "8.8 horas el 4 de diciembre de 2025",
    "D04": "(6 horas) fue el 2 de junio de 2026",
    "D05": "con 412 streams",
    "D06": "Bad Bunny ocupa el 6.2%",
    "D07": "«Tu No», 70",
    "D08": "«CI SARÒ», 74 streams",
    "D09": "ocupa el 5.3% de su tiempo",
    "D10": "francés (23 horas)",
    "D18": "«BAILE INoLVIDABLE» de Bad Bunny",
    "D19": "repite a Alfa («CI SARÒ»",
    "D11": "salta el 27% de las canciones de noche",
    "D12": "y el 16% en la mañana",
    "D13": "(31% contra 19%)",
    "D14": "(31% contra 19%)",
    "D15": "alrededor del 15%",
    "D16": "el 41% del tiempo de Gabriel",
    "D17": "el 42% del de Javier",
    "E01": "14 de julio de 2023",
    "E02": "(0.9 a 1)",
    "E03": "Gabriel escuchó el doble",
    "E06": "terminan el 3 (Javier) y el 25 (Gabriel)",
    "E04": "(10% de las reproducciones de",
    "E05": "vale 97 a 98%",
    "F01": "España, Reino Unido, Francia, Chequia, Alemania e Italia entre el 18 de junio "
    "y el 9 de julio de 2025",
    "F02": "entre el 30 de marzo y el 2 de junio de 2026",
    "F03": "Las 2,073 reproducciones",
    "F04": "desde noviembre de 2024",
    "E07": "empieza en noviembre de 2022, ocho meses antes",
}


def report_text() -> str:
    from docx import Document

    d = Document(INFORME)
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    return " ".join(" ".join(parts).split())  # espacios normalizados


@dataclass
class Check:
    id: str
    seccion: str
    afirmacion: str
    informe: str
    sql: str | Callable[[duckdb.DuckDBPyConnection], object]
    fmt: Callable[[object], str] = str


def n0(v) -> str:
    return f"{round(v):,}"


def p0(v) -> str:
    return f"{round(100 * v)}%"


def p1(v) -> str:
    return f"{100 * v:.1f}%"


def h1(v) -> str:
    return f"{v:.1f}"


def connect() -> duckdb.DuckDBPyConnection:
    c = duckdb.connect()
    c.execute("set TimeZone = 'UTC'")
    c.execute(f"create view a as select * from '{SILVER_PATH.as_posix()}'")
    c.execute(f"create view p as select * from a where ts_local >= '{W}'")  # ventana común
    # "música" = todo lo que no es episodio de podcast (igual que app.js)
    c.execute("create view m as select * from p where content_type <> 'episode'")
    return c


def by_user(expr: str, view: str = "p", where: str = "true") -> Callable:
    """Devuelve una función que calcula `expr` para un usuario: by_user(...)('gabo')."""

    def make(u: str) -> str:
        return f"select {expr} from {view} where user_id = '{u}' and ({where})"

    return make


# ---------- cálculos que no caben en una expresión ----------
def share(c, u, where: str, view: str = "p") -> float:
    return c.sql(
        f"select sum(ms_played) filter (where {where}) / sum(ms_played) from {view} "
        f"where user_id = '{u}'"
    ).fetchone()[0]


def day_share(c, u) -> tuple[int, float]:
    n, span = c.sql(
        f"""select count(distinct ts_local::date),
                  datediff('day', min(ts_local)::date, max(ts_local)::date) + 1
           from p where user_id = '{u}'"""
    ).fetchone()
    return n, n / span


def median_min(c, u) -> float:
    # mismo criterio que app.js: días ordenados, elemento floor(n/2) (base 0)
    return c.sql(
        f"""with d as (select ts_local::date d, sum(round(ms_played / 1000)) s from p
                       where user_id = '{u}' group by 1)
           select list_sort(list(s))[(count(*) // 2) + 1] / 60 from d"""
    ).fetchone()[0]


def peak_hour(c, u) -> int:
    return c.sql(
        f"""select hour(ts_local) h from p where user_id = '{u}' group by 1
           order by sum(ms_played) desc limit 1"""
    ).fetchone()[0]


def top10_share(c, u) -> float:
    return c.sql(
        f"""with x as (select artist_name, sum(ms_played) ms from m
                       where user_id = '{u}' and artist_name is not null group by 1)
           select (select sum(ms) from (select ms from x order by ms desc limit 10))
                  / (select sum(ms_played) from m where user_id = '{u}')"""
    ).fetchone()[0]


def main_device(c, u) -> str:
    plat, s = c.sql(
        f"""select platform_family, sum(ms_played) / (select sum(ms_played) from p
                                                       where user_id = '{u}')
           from p where user_id = '{u}' group by 1 order by 2 desc limit 1"""
    ).fetchone()
    return f"{ {'android': 'Android', 'windows': 'Windows'}.get(plat, plat) } {p0(s)}"


def monthly_avg(c, u, year: int) -> float:
    return c.sql(
        f"""select sum(ms_played) / 3.6e6 / count(distinct date_trunc('month', ts_local))
           from p where user_id = '{u}' and year(ts_local) = {year}"""
    ).fetchone()[0]


def quarter_leaders(c, u) -> list[str]:
    rows = c.sql(
        f"""select q, artist_name from (
              select q, artist_name, row_number() over (partition by q order by ms desc) rk
              from (select year(ts_local) * 10 + quarter(ts_local) q, artist_name,
                           sum(ms_played) ms
                    from m where user_id = '{u}' and artist_name is not null group by 1, 2))
           where rk = 1 order by q"""
    ).fetchall()
    return [r[1] for r in rows]


def streak(c, u, artist: str) -> int:
    n = 0
    for a in reversed(quarter_leaders(c, u)):
        if a != artist:
            break
        n += 1
    return n


def in_last(c, u, artist: str, k: int) -> int:
    return sum(a == artist for a in quarter_leaders(c, u)[-k:])


def first_arrival(c) -> tuple[int, int, int]:
    """Regla de app.js renderShared: artistas compartidos, primer stream de cada uno en la
    ventana, umbral = 1.er día del mes del inicio común + 3 meses, diferencia > 30 días."""
    rows = c.sql(
        """with f as (select user_id, artist_name, min(ts_local) t from m
                      where is_stream and artist_name is not null group by 1, 2),
                s as (select artist_name from (select distinct user_id, artist_name from m
                      where artist_name is not null) group by 1 having count(*) = 2),
                start as (select date_trunc('month', max(t0)) + interval 3 month th
                          from (select min(ts_local) t0 from p group by user_id))
           select (select count(*) from s),
                  count(*) filter (where g.t < j.t - interval 30 day),
                  count(*) filter (where j.t < g.t - interval 30 day)
           from s join f g on g.artist_name = s.artist_name and g.user_id = 'gabo'
                  join f j on j.artist_name = s.artist_name and j.user_id = 'javi', start
           where g.t >= start.th and j.t >= start.th"""
    ).fetchone()
    return rows


def shared_share(c, u) -> float:
    return c.sql(
        f"""with s as (select artist_name from (select distinct user_id, artist_name from m
                       where artist_name is not null) group by 1 having count(*) = 2)
           select sum(ms_played) filter (where artist_name in (select * from s))
                  / sum(ms_played) from m where user_id = '{u}'"""
    ).fetchone()[0]


def record_day(c, u) -> str:
    d, h = c.sql(
        f"""select ts_local::date, sum(ms_played) / 3.6e6 from p where user_id = '{u}'
           group by 1 order by 2 desc limit 1"""
    ).fetchone()
    return f"{h:.1f} h, {d}"


def top_cell(c, u) -> str:
    dow, h = c.sql(
        f"""select isodow(ts_local), hour(ts_local) from p where user_id = '{u}'
           group by 1, 2 order by sum(ms_played) desc limit 1"""
    ).fetchone()
    return f"{['lun', 'mar', 'mié', 'jue', 'vie', 'sáb', 'dom'][dow - 1]} {h} h"


def streams_of(c, u, track: str) -> int:
    return c.sql(
        f"""select count(*) from m where user_id = '{u}' and is_stream and track_name = ?""",
        params=[track],
    ).fetchone()[0]


def continuous_start(c) -> str:
    rows = c.sql(
        """select strftime(date_trunc('month', ts_local), '%Y-%m') m, count(*) n from a
           where user_id = 'javi' and ts_local < '2023-07-01' group by 1 order by 1"""
    ).fetchall()
    start = None
    for m, n in rows:
        start = (start or m) if n >= 50 else None
    return start


def top_song(c, u) -> str:
    # una canción = nombre + artista (las versiones single y de álbum tienen URIs distintas)
    name, n = c.sql(
        f"""select track_name, count(*) from m where user_id = '{u}' and is_stream
           and track_name is not null group by track_name, artist_name order by 2 desc limit 1"""
    ).fetchone()
    return f"{name} ({n})"


def skip_block(c, u, lo: int, hi: int) -> float:
    return c.sql(
        f"""select avg(is_skip::int) from m where user_id = '{u}'
           and hour(ts_local) >= {lo} and hour(ts_local) < {hi}"""
    ).fetchone()[0]


def skip_range(c, u) -> str:
    v = [skip_block(c, u, lo, lo + 6) for lo in (0, 6, 12, 18)]
    n = c.sql(
        f"""select count(*) filter (where hour(ts_local) < 6) from m where user_id = '{u}'"""
    ).fetchone()[0]
    # app.js ignora bloques con < 200 reproducciones (la madrugada de Javier)
    v = v[1:] if n < 200 else v
    return f"{round(100 * min(v))} a {round(100 * max(v))}%"


def trip(c, u, lo: str, hi: str) -> str:
    rows = c.sql(
        f"""select country, min(ts_local)::date, max(ts_local)::date from p
           where user_id = '{u}' and country <> 'MX' and ts_local between '{lo}' and '{hi}'
           group by 1 order by 2"""
    ).fetchall()
    first = min(r[1] for r in rows)
    last = max(r[2] for r in rows)
    return f"{','.join(r[0] for r in rows)} {first} a {last}"


def data_js(c) -> str:
    """El dashboard lee data.js: mismo número de filas y mismas horas que silver."""
    text = DATA_JS.read_text(encoding="utf-8")
    d = json.loads(re.sub(r"^window\.SPOTIFY=|;\s*$", "", text))
    rows = d["rows"]
    out = []
    for ui, u in enumerate(d["users"]):
        secs = sum(s for uu, s in zip(rows["u"], rows["s"], strict=True) if uu == ui)
        silver = c.sql(f"select sum(ms_played) / 1000 from a where user_id = '{u}'").fetchone()[0]
        out.append(f"{u} {abs(secs - silver) / silver:.4%}")
    return f"{len(rows['u']):,} filas; diferencia de horas: " + ", ".join(out)


# ---------- afirmaciones del informe ----------
def checks() -> list[Check]:
    g, j = "gabo", "javi"
    hours = by_user("sum(ms_played) / 3.6e6")
    artists = by_user("count(distinct artist_name)", "m")
    skip = by_user("avg(is_skip::int)", "m")
    pods = by_user("sum(ms_played) filter (where content_type = 'episode') / sum(ms_played)")
    raw = f"read_json('{RAW_GLOB}', format = 'array', union_by_name = true)"
    return [
        # Datos y pipeline
        Check("A01", "1, 2", "Reproducciones en silver", "81,418", "select count(*) from a", n0),
        Check(
            "A02",
            "2",
            "Reproducciones de Gabriel",
            "61,422",
            f"select count(*) from a where user_id = '{g}'",
            n0,
        ),
        Check(
            "A03",
            "2",
            "Reproducciones de Javier",
            "19,996",
            f"select count(*) from a where user_id = '{j}'",
            n0,
        ),
        Check(
            "A04",
            "2",
            "Duplicados exactos eliminados",
            "11",
            f"select (select count(*) from {raw}) - (select count(*) from a)",
            n0,
        ),
        Check(
            "A05",
            "2",
            "Campos por reproducción en el export",
            "23",
            f"select count(*) from (describe select * from {raw})",
            n0,
        ),
        Check(
            "A06",
            "5",
            "Reproducciones con IP en el export (≈81 mil)",
            "81,429",
            f"select count(ip_addr) from {raw}",
            n0,
        ),
        Check(
            "A07",
            "5",
            "Silver no tiene columna de IP",
            "0",
            "select count(*) from (describe a) where column_name = 'ip_addr'",
            n0,
        ),
        # Tabla del resumen
        Check("B01", "1", "Horas, Gabriel", "2,841", hours(g), n0),
        Check("B02", "1", "Horas, Javier", "713", hours(j), n0),
        Check(
            "B03",
            "1",
            "Gabriel escucha unas 4 veces más horas",
            "4.0",
            f"select ({hours(g)}) / ({hours(j)})",
            h1,
        ),
        Check(
            "B04",
            "1",
            "Días con escucha, Gabriel",
            "95% (1,087 días)",
            lambda c: day_share(c, g),
            lambda v: f"{p0(v[1])} ({v[0]:,} días)",
        ),
        Check(
            "B05",
            "1",
            "Días con escucha, Javier",
            "49% (546 días)",
            lambda c: day_share(c, j),
            lambda v: f"{p0(v[1])} ({v[0]:,} días)",
        ),
        Check(
            "B06",
            "1",
            "Mediana por día con escucha, Gabriel",
            "139 min",
            lambda c: median_min(c, g),
            lambda v: f"{round(v)} min",
        ),
        Check(
            "B07",
            "1",
            "Mediana por día con escucha, Javier",
            "63 min",
            lambda c: median_min(c, j),
            lambda v: f"{round(v)} min",
        ),
        Check("B08", "1", "Hora pico, Gabriel", "14", lambda c: peak_hour(c, g)),
        Check("B09", "1", "Hora pico, Javier", "18", lambda c: peak_hour(c, j)),
        Check(
            "B10",
            "1",
            "Segundo pico de Gabriel a las 6 h (% del tiempo a las 6 h)",
            "6.7%",
            f"select sum(ms_played) filter (where hour(ts_local) = 6) / sum(ms_played) "
            f"from p where user_id = '{g}'",
            p1,
        ),
        Check("B11", "1", "Artistas distintos, Gabriel", "1,662", artists(g), n0),
        Check("B12", "1", "Artistas distintos, Javier", "1,058", artists(j), n0),
        Check(
            "B13",
            "1",
            "Peso de sus 10 artistas principales, Gabriel",
            "21%",
            lambda c: top10_share(c, g),
            p0,
        ),
        Check(
            "B14",
            "1",
            "Peso de sus 10 artistas principales, Javier",
            "28%",
            lambda c: top10_share(c, j),
            p0,
        ),
        Check("B15", "1", "Canciones saltadas, Gabriel", "21%", skip(g), p0),
        Check("B16", "1", "Canciones saltadas, Javier", "15%", skip(j), p0),
        Check(
            "B17",
            "1, 3",
            "Dispositivo principal, Gabriel",
            "Android 92%",
            lambda c: main_device(c, g),
        ),
        Check(
            "B18",
            "1, 3",
            "Dispositivo principal, Javier",
            "Windows 70%",
            lambda c: main_device(c, j),
        ),
        Check("B19", "1", "Tiempo en podcasts, Gabriel", "0.1%", pods(g), p1),
        Check("B20", "1", "Tiempo en podcasts, Javier", "5.2%", pods(j), p1),
        # Viñetas del resumen
        Check(
            "C01",
            "1",
            "Javier: tiempo entre 17 y 21 h",
            "40%",
            lambda c: share(c, j, "hour(ts_local) >= 17 and hour(ts_local) < 21"),
            p0,
        ),
        Check(
            "C02",
            "1",
            "Javier: tiempo de madrugada (0 a 6 h)",
            "0.4%",
            lambda c: share(c, j, "hour(ts_local) < 6"),
            p1,
        ),
        Check(
            "C03",
            "1",
            "Gabriel: horas por mes en 2023",
            "42",
            lambda c: monthly_avg(c, g, 2023),
            n0,
        ),
        Check(
            "C04",
            "1",
            "Gabriel: horas por mes en 2025",
            "104",
            lambda c: monthly_avg(c, g, 2025),
            n0,
        ),
        Check(
            "C05", "1", "Javier: horas por mes en 2025", "10", lambda c: monthly_avg(c, j, 2025), n0
        ),
        Check(
            "C06",
            "1",
            "Bad Bunny, nº 1 de Gabriel: trimestres seguidos",
            "6",
            lambda c: streak(c, g, "Bad Bunny"),
        ),
        Check(
            "C07",
            "1",
            "Alfa, nº 1 de Javier en los últimos 9 trimestres",
            "8",
            lambda c: in_last(c, j, "Alfa", 9),
        ),
        Check(
            "C08",
            "1, 3",
            "Compartidos / Gabriel llegó primero / Javier llegó primero",
            "205 / 27 / 55",
            first_arrival,
            lambda v: f"{v[0]} / {v[1]} / {v[2]}",
        ),
        # Hallazgos
        Check(
            "D01",
            "3",
            "Gabriel: fin de semana, % del tiempo",
            "20%",
            lambda c: share(c, g, "isodow(ts_local) >= 6"),
            p0,
        ),
        Check(
            "D02", "3", "Javier: celda más fuerte del heatmap", "dom 18 h", lambda c: top_cell(c, j)
        ),
        Check("D03", "3", "Día récord, Gabriel", "8.8 h, 2025-12-04", lambda c: record_day(c, g)),
        Check("D04", "3", "Día récord, Javier", "6.0 h, 2026-06-02", lambda c: record_day(c, j)),
        Check(
            "D05",
            "3",
            "Streams de «BAILE INoLVIDABLE» (Gabriel)",
            "412",
            lambda c: streams_of(c, g, "BAILE INoLVIDABLE"),
            n0,
        ),
        Check(
            "D06",
            "3",
            "Bad Bunny, % del tiempo de Gabriel en música",
            "6.2%",
            lambda c: share(c, g, "artist_name = 'Bad Bunny'", "m"),
            p1,
        ),
        Check(
            "D07", "3", "Streams de «Tu No» (Javier)", "70", lambda c: streams_of(c, j, "Tu No"), n0
        ),
        Check(
            "D08",
            "3",
            "Streams de «CI SARÒ» (Javier)",
            "74",
            lambda c: streams_of(c, j, "CI SARÒ"),
            n0,
        ),
        Check(
            "D09",
            "3",
            "Alfa, % del tiempo de Javier en música",
            "5.3%",
            lambda c: share(c, j, "artist_name = 'Alfa'", "m"),
            p1,
        ),
        Check(
            "D10",
            "3",
            "Podcast principal de Javier (francés) en horas",
            "23",
            f"select sum(ms_played) / 3.6e6 from p where user_id = '{j}' "
            "and show_name = 'Meuf, je t''ai pas dit!'",
            n0,
        ),
        Check(
            "D18",
            "3",
            "Canción más repetida de Gabriel (nombre + artista)",
            "BAILE INoLVIDABLE (412)",
            lambda c: top_song(c, g),
        ),
        Check(
            "D19",
            "3",
            "Canción más repetida de Javier (nombre + artista)",
            "CI SARÒ (74)",
            lambda c: top_song(c, j),
        ),
        Check(
            "D11",
            "3",
            "Gabriel: skips de noche (18 a 24 h)",
            "27%",
            lambda c: skip_block(c, g, 18, 24),
            p0,
        ),
        Check(
            "D12",
            "3",
            "Gabriel: skips en la mañana (6 a 12 h)",
            "16%",
            lambda c: skip_block(c, g, 6, 12),
            p0,
        ),
        Check(
            "D13",
            "3",
            "Gabriel: skips sin shuffle",
            "31%",
            f"select avg(is_skip::int) from m where user_id = '{g}' and not shuffle",
            p0,
        ),
        Check(
            "D14",
            "3",
            "Gabriel: skips con shuffle",
            "19%",
            f"select avg(is_skip::int) from m where user_id = '{g}' and shuffle",
            p0,
        ),
        Check(
            "D15",
            "3",
            "Javier: skips por bloque del día (alrededor del 15%)",
            "13 a 16%",
            lambda c: skip_range(c, j),
        ),
        Check(
            "D16",
            "3",
            "Artistas compartidos: % del tiempo de Gabriel",
            "41%",
            lambda c: shared_share(c, g),
            p0,
        ),
        Check(
            "D17",
            "3",
            "Artistas compartidos: % del tiempo de Javier",
            "42%",
            lambda c: shared_share(c, j),
            p0,
        ),
        # EDA
        Check(
            "E01",
            "4",
            "Inicio de los datos continuos de Gabriel",
            "2023-07-14",
            f"select min(ts_local)::date from a where user_id = '{g}' and ts_local >= '2023-01-01'",
        ),
        Check(
            "E02",
            "4",
            "Horas Gabriel / Javier en todo 2023 (ventana desde enero)",
            "0.9",
            "select sum(ms_played) filter (where user_id = 'gabo') / "
            "sum(ms_played) filter (where user_id = 'javi') from a where year(ts_local) = 2023",
            h1,
        ),
        Check(
            "E03",
            "4",
            "Horas Gabriel / Javier en jul a dic 2023",
            "2.0",
            "select sum(ms_played) filter (where user_id = 'gabo') / "
            "sum(ms_played) filter (where user_id = 'javi') from a "
            "where ts_local >= '2023-07-01' and year(ts_local) = 2023",
            h1,
        ),
        Check(
            "E06",
            "3",
            "Último dato (Figura 2 omite agosto de 2026)",
            "gabo 2026-08-25, javi 2026-08-03",
            "select string_agg(user_id || ' ' || d, ', ' order by user_id) from "
            "(select user_id, max(ts_local)::date d from a group by 1)",
        ),
        Check(
            "E07",
            "3",
            "Inicio del historial continuo de Javier (primer mes de una racha sin huecos "
            "con 50 reproducciones o más al mes)",
            "2022-11",
            continuous_start,
        ),
        Check(
            "E04",
            "4",
            "Canciones de Gabriel saltadas en 0 a 5 s",
            "10%",
            f"select avg((ms_played < 5000)::int) from m where user_id = '{g}' "
            "and content_type = 'track'",
            p0,
        ),
        Check(
            "E05",
            "4",
            "% skipped en fwdbtn / backbtn / endplay / trackdone",
            "97 / 98 / 98 / 0",
            "select list(round(100 * s) order by r) from (select reason_end r, "
            "avg(is_skip::int) s from a where content_type = 'track' and reason_end in "
            "('backbtn', 'endplay', 'fwdbtn', 'trackdone') group by 1)",
            lambda v: " / ".join(str(int(x)) for x in (v[2], v[0], v[1], v[3])),
        ),
        # Privacidad
        Check(
            "F01",
            "5",
            "Viaje de Gabriel (países y fechas)",
            "ES,GB,FR,CZ,DE,IT 2025-06-18 a 2025-07-09",
            lambda c: trip(c, g, "2025-06-01", "2025-07-31"),
        ),
        Check(
            "F02",
            "5",
            "Javier conectado desde Estados Unidos",
            "US 2026-03-30 a 2026-06-02",
            lambda c: trip(c, j, "2026-03-01", "2026-06-10"),
        ),
        Check(
            "F03",
            "5",
            "Reproducciones de Gabriel en incógnito",
            "2,073",
            f"select count(*) filter (where incognito) from a where user_id = '{g}'",
            n0,
        ),
        Check(
            "F04",
            "5",
            "Primera reproducción en incógnito",
            "2024-11",
            "select strftime(min(ts_local), '%Y-%m') from a where incognito",
        ),
        # Dashboard
        Check(
            "G01",
            "dashboard",
            "data.js contiene todas las filas y las mismas horas",
            "81,418 filas; diferencia de horas: gabo 0.0000%, javi 0.0000%",
            data_js,
        ),
    ]


def evaluate(c: duckdb.DuckDBPyConnection, ch: Check) -> str:
    v = ch.sql(c) if callable(ch.sql) else c.sql(ch.sql).fetchone()[0]
    return ch.fmt(v)


def main() -> int:
    c = connect()
    text = report_text()
    rows, bad = [], 0
    for ch in checks():
        got = evaluate(c, ch)
        ok = got == ch.informe
        # data.js redondea a segundos: se acepta una diferencia < 0.01 %
        if ch.id == "G01":
            ok = all(float(x.rstrip("%")) < 0.01 for x in re.findall(r"([\d.]+%)", got))
        cita = CITAS.get(ch.id, "")
        found = not cita or cita in text
        bad += not (ok and found)
        estado = "OK" if ok and found else "**REVISAR cifra**" if not ok else "**REVISAR texto**"
        rows.append(
            f"| {ch.id} | {ch.seccion} | {ch.afirmacion} | {cita or '(dashboard)'} | "
            f"{ch.informe} | {got} | {estado} |"
        )
    header = [
        "# Auditoría de cifras del informe",
        "",
        "Generado por `uv run python -m pipeline.auditoria` (también corre dentro de "
        "`generar.cmd`). No editar a mano. Método en [tecnico.md](tecnico.md) §9.",
        "",
        "- **Cita**: texto literal que se busca en `out/informe.docx`.",
        "- **Esperado**: el valor que afirma la cita, con el redondeo del informe.",
        "- **Calculado**: el mismo valor recalculado desde `data/silver/plays.parquet`.",
        "- **OK**: la cita está en el informe y el valor calculado es idéntico.",
        "",
        f"Ventana común: desde {W}. Resultado: **{len(rows) - bad} de {len(rows)} OK**.",
        "",
        "| ID | § | Afirmación | Cita en el informe | Esperado | Calculado | Estado |",
        "|---|---|---|---|---|---|---|",
    ]
    OUT.write_text("\n".join(header + rows) + "\n", encoding="utf-8")
    print(f"{OUT.relative_to(ROOT)}: {len(rows) - bad}/{len(rows)} OK")
    for r in rows:
        if "REVISAR" in r:
            print(r)
    return bad


if __name__ == "__main__":
    raise SystemExit(main())
